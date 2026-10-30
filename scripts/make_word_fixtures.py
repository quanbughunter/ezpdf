#!/usr/bin/env python3
"""
Word fixtures for the .docx reader/writer tests (tests/word-probe.ts,
scripts/e2e-word.cjs). Three producers, so the reader meets three styles of
WordprocessingML:

  word-a.docx       python-docx (Word's own default template): headings,
                    mixed run formatting, Vietnamese text, bullet and
                    numbered lists, a table with merged cells, a picture, a
                    hyperlink, a bookmark, a tab, a line break, a page break,
                    a PAGE field, a header and a footer, a paragraph with a
                    tracked change, an equation, a content control
  word-b.docx       python-docx, the hard cases for saving: a section break
                    (portrait → landscape), a heading wrapped in a _Toc
                    bookmark, a comment, a table cell ending in a bookmark's
                    end, a run with a Word 2010 text effect (w14), a field
                    in the middle of a sentence, a paragraph whose mark is a
                    tracked insertion
  word-lo.docx      word-a.docx saved again by LibreOffice
  word-pandoc.docx  written by pandoc from Markdown

Usage: python3 scripts/make_word_fixtures.py   (needs python-docx, Pillow;
LibreOffice and pandoc for the other two)
"""
import io
import os
import shutil
import subprocess
import tempfile

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image, ImageDraw

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, 'tests', 'fixtures')


def picture():
    img = Image.new('RGB', (240, 120), (11, 110, 105))
    d = ImageDraw.Draw(img)
    d.rectangle([12, 12, 228, 108], outline=(255, 255, 255), width=4)
    d.text((40, 50), 'ezPDF', fill=(255, 255, 255))
    buf = io.BytesIO()
    img.save(buf, 'PNG')
    buf.seek(0)
    return buf


def add_hyperlink(paragraph, url, text):
    part = paragraph.part
    r_id = part.relate_to(url, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
    link = OxmlElement('w:hyperlink')
    link.set(qn('r:id'), r_id)
    run = OxmlElement('w:r')
    rpr = OxmlElement('w:rPr')
    style = OxmlElement('w:rStyle')
    style.set(qn('w:val'), 'Hyperlink')
    rpr.append(style)
    color = OxmlElement('w:color')
    color.set(qn('w:val'), '0563C1')
    rpr.append(color)
    u = OxmlElement('w:u')
    u.set(qn('w:val'), 'single')
    rpr.append(u)
    run.append(rpr)
    t = OxmlElement('w:t')
    t.text = text
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


def add_field(paragraph, instr, result):
    def fld(kind):
        r = OxmlElement('w:r')
        f = OxmlElement('w:fldChar')
        f.set(qn('w:fldCharType'), kind)
        r.append(f)
        return r
    paragraph._p.append(fld('begin'))
    r = OxmlElement('w:r')
    it = OxmlElement('w:instrText')
    it.set(qn('xml:space'), 'preserve')
    it.text = f' {instr} '
    r.append(it)
    paragraph._p.append(r)
    paragraph._p.append(fld('separate'))
    paragraph.add_run(result)
    paragraph._p.append(fld('end'))


def word_a(path):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.5)
    sec.top_margin = sec.bottom_margin = Cm(2)
    sec.header.paragraphs[0].text = 'Công ty ezPDF · Đầu trang'
    fp = sec.footer.paragraphs[0]
    fp.text = 'Trang '
    add_field(fp, 'PAGE', '1')

    doc.add_heading('Hợp đồng dịch vụ', level=1)

    p = doc.add_paragraph()
    p.add_run('Đoạn mở đầu có ')
    r = p.add_run('chữ đậm')
    r.bold = True
    p.add_run(', ')
    r = p.add_run('chữ nghiêng')
    r.italic = True
    p.add_run(', ')
    r = p.add_run('gạch chân')
    r.underline = True
    p.add_run(' và ')
    r = p.add_run('chữ đỏ cỡ 14')
    r.font.color.rgb = RGBColor(0xB4, 0x23, 0x18)
    r.font.size = Pt(14)
    p.add_run('. Tiếng Việt đủ dấu: ắ ằ ẳ ẵ ặ ơ ư đ.')

    p = doc.add_paragraph('Đoạn thứ hai căn giữa, phông Times New Roman.')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in p.runs:
        run.font.name = 'Times New Roman'

    p = doc.add_paragraph('Đoạn để xóa: dòng này sẽ bị xóa trong bài kiểm.')

    doc.add_heading('Danh sách', level=2)
    doc.add_paragraph('Mục gạch đầu dòng một', style='List Bullet')
    doc.add_paragraph('Mục gạch đầu dòng hai', style='List Bullet')
    doc.add_paragraph('Bước một', style='List Number')
    doc.add_paragraph('Bước hai', style='List Number')
    doc.add_paragraph('Bước ba', style='List Number')

    doc.add_heading('Bảng', level=2)
    t = doc.add_table(rows=3, cols=3)
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(t.rows):
        for j, cell in enumerate(row.cells):
            cell.text = f'Ô {i + 1}.{j + 1}'
    a = t.cell(0, 0).merge(t.cell(0, 1))  # horizontal merge
    a.text = 'Ô gộp ngang'
    b = t.cell(1, 2).merge(t.cell(2, 2))  # vertical merge
    b.text = 'Ô gộp dọc'

    p = doc.add_paragraph('Hình: ')
    p.add_run().add_picture(picture(), width=Cm(4))

    p = doc.add_paragraph('Liên kết: ')
    add_hyperlink(p, 'https://example.com/ezpdf', 'trang ví dụ')
    p.add_run(' ở giữa câu.')

    # a bookmark around a whole paragraph
    p = doc.add_paragraph()
    start = parse_xml(f'<w:bookmarkStart {nsdecls("w")} w:id="7" w:name="Muc_tieu"/>')
    end = parse_xml(f'<w:bookmarkEnd {nsdecls("w")} w:id="7"/>')
    p._p.append(start)
    p.add_run('Đoạn có dấu trang (bookmark).')
    p._p.append(end)

    p = doc.add_paragraph('Cột\tsau tab')
    p = doc.add_paragraph('Dòng một')
    p.runs[0].add_break()
    p.add_run('dòng hai sau ngắt dòng')

    p = doc.add_paragraph('Trang hiện tại: ')
    add_field(p, 'PAGE', '1')

    # tracked change: read-only in ezPDF
    p = doc.add_paragraph('Đoạn có ')
    ins = parse_xml(
        f'<w:ins {nsdecls("w")} w:id="1" w:author="Kiểm thử" w:date="2026-10-30T10:00:00Z">'
        '<w:r><w:t xml:space="preserve">chữ chèn thêm </w:t></w:r></w:ins>'
    )
    p._p.append(ins)
    dele = parse_xml(
        f'<w:del {nsdecls("w")} w:id="2" w:author="Kiểm thử" w:date="2026-10-30T10:00:00Z">'
        '<w:r><w:delText>chữ bị xóa</w:delText></w:r></w:del>'
    )
    p._p.append(dele)
    p.add_run(' theo dõi thay đổi.')

    # an equation: read-only
    p = doc.add_paragraph('Công thức: ')
    math = parse_xml(
        '<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">'
        '<m:r><m:t>E=mc²</m:t></m:r></m:oMath>'
    )
    p._p.append(math)

    # a block content control
    sdt = parse_xml(
        f'<w:sdt {nsdecls("w")}><w:sdtPr><w:alias w:val="Ô nội dung"/></w:sdtPr><w:sdtContent>'
        '<w:p><w:r><w:t>Nội dung trong ô điều khiển.</w:t></w:r></w:p></w:sdtContent></w:sdt>'
    )
    doc.element.body.insert(len(doc.element.body) - 1, sdt)

    p = doc.add_paragraph('Trước ngắt trang.')
    p.runs[0].add_break(WD_BREAK.PAGE)
    p.add_run('Sau ngắt trang.')

    doc.add_paragraph('Đoạn cuối của tài liệu.')
    doc.save(path)


def word_b(path):
    from docx.enum.section import WD_ORIENT, WD_SECTION
    doc = Document()
    # a heading wrapped in a bookmark, as Word does for a table of contents
    h = doc.add_heading('', level=1)
    h._p.append(parse_xml(f'<w:bookmarkStart {nsdecls("w")} w:id="11" w:name="_Toc1"/>'))
    h.add_run('Chương một')
    h._p.append(parse_xml(f'<w:bookmarkEnd {nsdecls("w")} w:id="11"/>'))
    p = doc.add_paragraph('Câu có ')
    r = p.add_run('lời bình')
    p.add_run(' ở giữa.')
    doc.add_comment(r, text='Xem lại chỗ này', author='Kiểm thử', initials='KT')
    p = doc.add_paragraph('Số trang: ')
    add_field(p, 'PAGE', '1')
    p.add_run(' là trường giữa câu.')
    p = doc.add_paragraph()
    r = p.add_run('Chữ có viền (hiệu ứng Word 2010)')
    rpr = r._r.get_or_add_rPr()
    rpr.append(parse_xml(
        '<w14:textOutline xmlns:w14="http://schemas.microsoft.com/office/word/2010/wordml" w14:w="9525" w14:cap="flat" w14:cmpd="sng" w14:algn="ctr">'
        '<w14:solidFill><w14:srgbClr w14:val="0B6E69"/></w14:solidFill></w14:textOutline>'
    ))
    t = doc.add_table(rows=2, cols=2)
    t.style = 'Table Grid'
    t.cell(0, 0).text = 'Ô sẽ sửa'
    t.cell(0, 1).text = 'Ô có dấu trang'
    tc = t.cell(0, 1)._tc
    tc.append(parse_xml(f'<w:bookmarkEnd {nsdecls("w")} w:id="21"/>'))
    tc.findall('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}p')[0].insert(0, parse_xml(f'<w:bookmarkStart {nsdecls("w")} w:id="21" w:name="O_dau_trang"/>'))
    t.cell(1, 0).text = 'Ô 2.1'
    t.cell(1, 1).text = 'Ô 2.2'
    # the paragraph mark itself is a tracked insertion: read-only
    p = doc.add_paragraph('Dấu đoạn được chèn khi theo dõi thay đổi.')
    p._p.get_or_add_pPr().append(parse_xml(
        f'<w:rPr {nsdecls("w")}><w:ins w:id="31" w:author="Kiểm thử" w:date="2026-10-30T10:00:00Z"/></w:rPr>'
    ))
    doc.add_paragraph('Cuối phần dọc.')
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = sec.page_height, sec.page_width
    doc.add_paragraph('Phần ngang.')
    doc.save(path)


def word_lo(src, path):
    exe = shutil.which('soffice')
    if not exe:
        print('skip word-lo.docx: no LibreOffice')
        return
    with tempfile.TemporaryDirectory() as tmp:
        shutil.copy(src, os.path.join(tmp, 'in.docx'))
        subprocess.run([exe, '--headless', '--convert-to', 'docx:MS Word 2007 XML', '--outdir', os.path.join(tmp, 'out'), os.path.join(tmp, 'in.docx')],
                       check=True, capture_output=True, timeout=180)
        shutil.copy(os.path.join(tmp, 'out', 'in.docx'), path)


MD = """# Báo cáo tuần

Đoạn **đậm**, *nghiêng* và [liên kết](https://example.com).

- Việc một
- Việc hai
    - Việc con

1. Thứ nhất
2. Thứ hai

| Cột A | Cột B |
|-------|-------|
| một   | hai   |
| ba    | bốn   |

> Trích dẫn ngắn.

Đoạn cuối.
"""


def word_pandoc(path):
    exe = shutil.which('pandoc')
    if not exe:
        print('skip word-pandoc.docx: no pandoc')
        return
    subprocess.run([exe, '-f', 'markdown', '-t', 'docx', '-o', path], input=MD.encode(), check=True, timeout=120)


if __name__ == '__main__':
    a = os.path.join(OUT, 'word-a.docx')
    word_a(a)
    word_b(os.path.join(OUT, 'word-b.docx'))
    word_lo(a, os.path.join(OUT, 'word-lo.docx'))
    word_pandoc(os.path.join(OUT, 'word-pandoc.docx'))
    for f in ('word-a.docx', 'word-b.docx', 'word-lo.docx', 'word-pandoc.docx'):
        p = os.path.join(OUT, f)
        if os.path.exists(p):
            print(f, os.path.getsize(p), 'bytes')
