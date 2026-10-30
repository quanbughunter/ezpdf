#!/usr/bin/env python3
"""
Reads a .docx saved by ezPDF with python-docx (an independent reader) and
checks it against expectations written by tests/word-probe.ts.
Prints {"checks": [{"name", "ok", "detail"}]} as JSON.
"""
import json
import sys

from docx import Document
from docx.oxml.ns import qn

if sys.argv[1] == '--drawings':
    # how many drawings (wp:docPr) and how many different ids
    import re
    import zipfile
    x = zipfile.ZipFile(sys.argv[2]).read('word/document.xml').decode('utf-8')
    ids = re.findall(r'<wp:docPr[^>]*?\sid="(\d+)"', x)
    print(len(ids), len(set(ids)))
    sys.exit(0)

path, expect_path = sys.argv[1], sys.argv[2]
exp = json.load(open(expect_path, encoding='utf-8'))
doc = Document(path)
checks = []


def check(name, ok, detail=''):
    checks.append({'name': name, 'ok': bool(ok), 'detail': detail})


paras = doc.paragraphs
alltext = '\n'.join(p.text for p in paras)


def runs_with(text):
    out = []
    for p in paras:
        for r in p.runs:
            if text in r.text or (r.text and r.text in text and len(r.text) > 3):
                out.append(r)
    return out


def run_prop(text, attr):
    rs = [r for p in paras for r in p.runs if r.text and r.text.strip() and r.text.strip() in text]
    # the effective value: direct, else from the character style, else the paragraph style
    vals = []
    for r in rs:
        v = getattr(r, attr)
        if v is None and r.style is not None:
            v = getattr(r.style.font, attr)
        if v is None:
            for p in paras:
                if r._r in list(p._p.iter()):
                    v = getattr(p.style.font, attr)
        vals.append(bool(v))
    return vals


check('opens; text found', all(t in alltext for t in exp['texts']), ', '.join(t for t in exp['texts'] if t not in alltext))
check('deleted paragraph gone', all(t not in alltext for t in exp['absent']))
for t in exp['bold']:
    v = run_prop(t, 'bold')
    check(f'bold: “{t}”', v and all(v), str(v))
for t in exp['notItalic']:
    v = run_prop(t, 'italic')
    check(f'not italic: “{t}”', v and not any(v), str(v))
for t in exp['notUnderlined']:
    v = run_prop(t, 'underline')
    check(f'not underlined: “{t}”', v and not any(v), str(v))
r, c, text = exp['cell']
cell = doc.tables[0].cell(r, c).text
check('table cell edited, table still 3×3', cell == text and len(doc.tables[0].rows) == 3 and len(doc.tables[0].columns) == 3, cell)
for t in exp['centered']:
    ps = [p for p in paras if p.text == t]
    ok = ps and all((p.alignment is not None and p.alignment == 1) for p in ps)
    check(f'centred (kept from the split paragraph): “{t.strip()}”', ok, str([p.alignment for p in ps]))
for t in exp['times']:
    ps = [p for p in paras if p.text == t]
    ok = ps and all(r.font.name == 'Times New Roman' for p in ps for r in p.runs)
    check(f'font kept: “{t.strip()}” in Times New Roman', ok)
for t, style in exp['style'].items():
    ps = [p for p in paras if p.text == t]
    check(f'style of “{t}” is {style}', ps and ps[0].style.name == style, ps and ps[0].style.name or 'missing')
for t in exp['noNumbering']:
    ps = [p for p in paras if p.text == t]
    num = ps and ps[0]._p.pPr is not None and ps[0]._p.pPr.find(qn('w:numPr'))
    numid = num is not None and num is not False and num.find(qn('w:numId'))
    ok = ps and numid is not None and numid is not False and numid.get(qn('w:val')) == '0'
    check(f'“{t}” is out of the list (numId 0)', ok)
if exp.get('tracked'):
    body = doc.element.body
    check('tracked change kept (w:ins, w:del)', body.find('.//' + qn('w:ins')) is not None and body.find('.//' + qn('w:del')) is not None)
    check('equation kept', body.find('.//{http://schemas.openxmlformats.org/officeDocument/2006/math}oMath') is not None)
    check('picture kept', body.find('.//{http://schemas.openxmlformats.org/drawingml/2006/main}blip') is not None)
    check('hyperlink kept', body.find('.//' + qn('w:hyperlink')) is not None)
    check('bookmark kept', body.find('.//' + qn('w:bookmarkStart')) is not None and body.find('.//' + qn('w:bookmarkEnd')) is not None)
    check('field kept (PAGE)', any('PAGE' in (e.text or '') for e in body.iter(qn('w:instrText'))))

print(json.dumps({'checks': checks}, ensure_ascii=False))
